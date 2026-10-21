"""Build the QuantumBit Revenue Cloud + Field Service demo story deck from the
Salesforce FY27 corporate template (layouts/masters only)."""
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

TEMPLATE = sys.argv[1]
OUTPUT = sys.argv[2]

NAVY = RGBColor(0x03, 0x2D, 0x60)
BLUE = RGBColor(0x06, 0x6A, 0xFE)
DEEP = RGBColor(0x02, 0x2A, 0xC0)
LIGHT = RGBColor(0xEA, 0xF5, 0xFE)
PANEL = RGBColor(0xF3, 0xF8, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x6B, 0x85)
ORANGE_BG = RGBColor(0xFF, 0xEB, 0xD6)
ORANGE_TX = RGBColor(0xA8, 0x4B, 0x00)
ORANGE_LN = RGBColor(0xF3, 0x83, 0x04)

HEAD = "Avant Garde Demi SFDC"
BODY = "Salesforce Sans"

L_TITLE, L_CONTENT, L_SCENE, L_SECTION, L_THANKS = 0, 4, 16, 35, 56

# ---------------------------------------------------------------- helpers


def strip_slides(prs):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)


def remove_ph(slide, *idxs):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx in idxs:
            ph._element.getparent().remove(ph._element)


def set_ph(slide, idx, text, size=None):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text_frame.text = text
            for para in ph.text_frame.paragraphs:
                pPr = para._p.get_or_add_pPr()
                pPr.set("marL", "0")
                pPr.set("indent", "0")
                if size:
                    for r in para.runs:
                        r.font.size = Pt(size)
            return ph
    raise KeyError(idx)


def style_run(run, size, color, bold=False, font=BODY, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def bullet(paragraph, char="•", indent_in=0.3):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT,
         margin=0.0):
    """paras: list of dicts {t, size, color, bold, font, bullet, space}"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        if p.get("space"):
            para.space_before = Pt(p["space"])
        runs = p["t"] if isinstance(p["t"], list) else [(p["t"], {})]
        for chunk, over in runs:
            r = para.add_run()
            r.text = chunk
            style_run(r, over.get("size", p.get("size", 18)),
                      over.get("color", p.get("color", NAVY)),
                      over.get("bold", p.get("bold", False)),
                      over.get("font", p.get("font", BODY)),
                      over.get("italic", p.get("italic", False)))
        if p.get("bullet"):
            bullet(para, p.get("char", "•"))
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        dash=False, line_w=1.25):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4  # dash
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def pill(slide, x, y, label, vision=False, w=None, size=12):
    w = w or (0.22 + 0.105 * len(label))
    fill, color, line = (ORANGE_BG, ORANGE_TX, ORANGE_LN) if vision else (BLUE, WHITE, None)
    s = box(slide, x, y, w, 0.36, fill=fill, line=line, radius=0.5, line_w=1)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    style_run(r, size, color, bold=True)
    return s


def circle_num(slide, x, y, n, d=0.55, fill=BLUE, size=16):
    s = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    style_run(r, size, WHITE, bold=True, font=HEAD)
    return s


def arrow(slide, x1, y1, x2, y2, color=BLUE, w=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


# ---------------------------------------------------------------- content

SCENES = [
    dict(
        n=1, act=1, vision=False, title="Configure & quote",
        persona="Seller", app="Revenue Cloud · Quote & Configurator",
        sees=[
            "A seller builds the Infinitech quote: PowerSwerve R750 server bundle",
            "A product rule adds Installation & Assembly (QB-INST-1000) with the hardware",
            "Gold Hardware Maintenance is priced automatically at 20% of hardware",
        ],
        moment="Service is sold with the hardware, so the install is never forgotten.",
        shot="Product Configurator with the R750 bundle, then the quote line editor with pricing",
        path=["Open Infinitech opportunity → New Quote", "Add PowerSwerve R750 → Configure",
              "Show auto-added Installation & Assembly", "Add Gold Maintenance → show price waterfall"],
        talk="""Start with the seller. Infinitech is expanding its data center and wants QuantumBit servers.
Configure the PowerSwerve R750 bundle and point out that the installation service is added by a product rule — the seller doesn't have to remember it.
Add Gold Hardware Maintenance and show that pricing calculates it as 20% of hardware. Every line we sell here becomes work for Field Service later.""",
    ),
    dict(
        n=2, act=1, vision=False, title="Activate the order",
        persona="Fulfillment Manager", app="Revenue Cloud · Dynamic Revenue Orchestration",
        sees=[
            "The quote becomes an activated order with one click",
            "Dynamic Revenue Orchestration breaks the order into a fulfillment plan",
            "Ship → Field Installation → Installation Complete → Initiate Billing",
        ],
        moment="Billing waits on the install. The plan won't bill until the field says it's done.",
        shot="Fulfillment plan for the Infinitech order with the Field Installation step highlighted",
        path=["Quote → Create Order → Activate", "Open the order's fulfillment plan",
              "Walk the steps left to right", "Point at Initiate Billing waiting on Installation Complete"],
        talk="""Activate the order and switch to the fulfillment manager. Dynamic Revenue Orchestration decomposed the order into a plan.
Walk the steps: ship hardware, field installation, installation complete, then billing.
The key point: Initiate Billing depends on Installation Complete, so the customer is never billed for equipment that isn't running.""",
    ),
    dict(
        n=3, act=1, vision=False, title="Automatic dispatch",
        persona="Dispatcher", app="Field Service · Work Orders & Scheduling",
        sees=[
            "The Field Installation step creates a Work Order and Service Appointment automatically",
            "The Work Order links back to the Order and the fulfillment step",
            "The appointment lands in the San Francisco territory, ready to schedule",
        ],
        moment="No re-keying. The order became field work the moment it was activated.",
        shot="Work Order record showing the Order and Fulfillment Step links, plus its Service Appointment",
        path=["Open the Field Installation step → related Work Order", "Show Order and Fulfillment Step fields",
              "Open the Service Appointment", "Show territory and work type (Server Rack Install)"],
        talk="""Switch to the dispatcher. Nobody typed this Work Order — the Field Installation step created it, along with a Service Appointment.
Show the Order and Fulfillment Step links on the Work Order. That link is what lets the field close the loop back to Revenue Cloud later.
The work type and territory came from the Field Service setup, so it's ready to schedule.""",
    ),
    dict(
        n=4, act=1, vision=False, title="Technician on site",
        persona="Technician", app="Field Service · Mobile",
        sees=[
            "The job goes to a technician with the Server Rack Install skill",
            "On site, the technician records the server serial number on the Asset",
            "The technician completes the Work Order",
        ],
        moment="The person in the field closes the deal's last mile from a phone.",
        shot="Field Service mobile app: appointment details, Asset serial capture, Complete status",
        path=["Assign the appointment to the certified technician", "Open it on mobile (or the Work Order on desktop)",
              "Enter the serial number on the Asset", "Set the Work Order to Completed"],
        talk="""The appointment goes to a technician who holds the Server Rack Install skill.
On site, they rack the server, record its serial number on the same Asset record Revenue Cloud created, and complete the job.
That single status change is about to start billing.""",
    ),
    dict(
        n=5, act=1, vision=False, title="Billing starts at install",
        persona="Billing Specialist", app="Revenue Cloud · Billing",
        sees=[
            "Completing the Work Order completes the fulfillment step automatically",
            "Installation Complete releases Initiate Billing",
            "The billing schedule starts on the install date, not the order date",
        ],
        moment="Revenue starts when value starts. No credits, no disputes.",
        shot="Fulfillment plan with Installation Complete done, then the billing schedule start date",
        path=["Return to the fulfillment plan → Installation Complete is done", "Initiate Billing is now running",
              "Open the billing schedule", "Compare start date to the order date"],
        talk="""Go back to the fulfillment plan. Installation Complete flipped to done on its own when the technician completed the Work Order.
Billing is released, and the billing schedule starts on the install date.
This is the moment finance cares about: no invoicing for equipment that isn't live, so no credits or disputes.""",
    ),
    dict(
        n=6, act=2, vision=True, title="Maintenance runs itself",
        persona="Service Manager", app="Field Service · Maintenance Plans",
        sees=[
            "Gold maintenance on the order creates a Maintenance Plan on the new Asset",
            "Monthly preventive visits are generated as Work Orders automatically",
            "Each visit is tied to the Asset Revenue Cloud sold",
        ],
        moment="What was promised on the quote is scheduled in the field, automatically.",
        shot="Maintenance Plan on the R750 Asset with upcoming preventive Work Orders",
        path=["Open the R750 Asset", "Show the Maintenance Plan created from Gold maintenance",
              "Show the next generated preventive Work Orders"],
        talk="""Vision scene. Gold maintenance promised monthly on-site check-ups. Here that promise becomes a Maintenance Plan on the Asset, generating the visits automatically.
Nobody in service has to read the contract to know what's owed.""",
    ),
    dict(
        n=7, act=2, vision=True, title="Break/fix with coverage check",
        persona="Dispatcher & Billing", app="Field Service + Revenue Cloud Billing",
        sees=[
            "A power supply fails and a break/fix Work Order is opened",
            "Covered by Gold maintenance: no charge to the customer",
            "Not covered: parts and labor become a priced, invoiced order",
        ],
        moment="One record decides what the customer pays. Service pays for itself.",
        shot="Side by side: covered Work Order (no charge) vs uncovered Work Order with its billable order and invoice",
        path=["Open break/fix Work Order on Infinitech → Covered", "Repeat on an account without maintenance",
              "Show the billable order created from parts and labor", "Show the invoice in Billing"],
        talk="""Vision scene. Same failure, two customers. Infinitech has Gold maintenance, so the coverage check says no charge.
A customer without maintenance gets a billable order for parts and labor, priced by Revenue Cloud and invoiced by Billing.
This is the strongest moment of the demo: entitlement and revenue decided by the same record.""",
    ),
    dict(
        n=8, act=2, vision=True, title="Upsell from the field",
        persona="Technician → Seller", app="Field Service Mobile + Revenue Cloud",
        sees=[
            "On site, the technician sees the rack is nearly full",
            "One quick action from the Work Order flags an expansion opportunity",
            "Sales amends the existing asset instead of starting a new deal",
        ],
        moment="Growth starts with the person who's actually in the data center.",
        shot="Quick action on the Work Order, then the amendment quote on the existing Asset",
        path=["Work Order → Flag Expansion quick action", "Opportunity appears for the account owner",
              "Seller amends the R750 Asset → adds capacity"],
        talk="""Vision scene. The technician is the only person who's physically seen the customer's rack. One tap flags an expansion opportunity.
The seller amends the existing Asset rather than quoting from scratch, so the new lines co-term with what's already installed.""",
    ),
    dict(
        n=9, act=2, vision=False, title="One install base",
        persona="Account Team", app="Revenue Cloud + Field Service · Asset",
        sees=[
            "Revenue Cloud and Field Service work from the same Asset record",
            "Orders, Work Orders, and billing all appear on one Asset",
            "Renewals and amendments quote against what's actually installed",
        ],
        moment="Sales, service, and finance finally agree on what the customer has.",
        shot="R750 Asset record with related Order, Work Orders, and billing together",
        path=["Open the R750 Asset", "Show related Order and fulfillment history",
              "Show related Work Orders and serial number", "Show billing details"],
        talk="""Close on the Asset. This one record carries the order, the install Work Order, the serial number, and billing.
There's no integration between two systems because there's one system. When renewal comes, the seller quotes what's really in the rack.""",
    ),
]

# ---------------------------------------------------------------- build

prs = Presentation(TEMPLATE)
strip_slides(prs)
M = prs.slide_masters[1]


def layout(i):
    return M.slide_layouts[i]


# 1. Title
s = prs.slides.add_slide(layout(L_TITLE))
set_ph(s, 0, "From quote to installed, and back again")
set_ph(s, 1, "Revenue Cloud + Field Service: the QuantumBit demo story")
set_ph(s, 2, "Demo storyboard · Draft")
notes(s, "Purpose: walk the audience through one customer lifecycle that crosses Revenue Cloud and Field Service in both directions. Every slide after this maps to something we show live.")

# 2. Scenario
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "Meet Infinitech")
set_ph(s, 1, "Expanding its San Francisco data center with QuantumBit")
remove_ph(s, 2, 3)

text(s, 0.94, 3.05, 9, 0.4, [dict(t="WHAT THEY'RE BUYING", size=14, color=BLUE, bold=True)])
products = [
    ("PowerSwerve R750 Rack Server", "QB-BDL-R750", "The hardware Field Service installs and maintains", "Scenes 1, 4, 9"),
    ("Installation and Assembly", "QB-INST-1000", "Becomes the install Work Order", "Scenes 2–5"),
    ("Gold Hardware Maintenance", "QB-SUPP-1005", "Monthly on-site check-ups and break/fix coverage", "Scenes 6–7"),
]
for i, (name, sku, why, where) in enumerate(products):
    y = 3.55 + i * 1.95
    box(s, 0.94, y, 8.6, 1.7)
    text(s, 1.25, y + 0.2, 6.0, 0.5, [dict(t=name, size=21, bold=True, font=HEAD)])
    text(s, 6.9, y + 0.24, 2.4, 0.4, [dict(t=sku, size=13, color=MUTED, align=PP_ALIGN.RIGHT)])
    text(s, 1.25, y + 0.75, 8.0, 0.45, [dict(t=why, size=17, color=NAVY)])
    text(s, 1.25, y + 1.18, 8.0, 0.4, [dict(t=where, size=13, color=BLUE, bold=True)])

text(s, 10.3, 3.05, 8.7, 0.4, [dict(t="THE CAST", size=14, color=BLUE, bold=True)])
cast = [
    ("Seller", "Quotes and configures", "Revenue Cloud"),
    ("Fulfillment Manager", "Watches the orchestration plan", "Revenue Cloud"),
    ("Dispatcher", "Schedules the field work", "Field Service"),
    ("Technician", "Installs, services, spots growth", "Field Service Mobile"),
    ("Billing Specialist", "Starts billing at go-live", "Revenue Cloud Billing"),
]
for i, (role, does, app) in enumerate(cast):
    y = 3.55 + i * 1.18
    box(s, 10.3, y, 8.76, 1.0)
    circle_num(s, 10.5, y + 0.2, role[0], d=0.6, fill=DEEP if "Field" in app else BLUE, size=17)
    text(s, 11.35, y + 0.05, 3.9, 0.9, [dict(t=role, size=19, bold=True)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 14.05, y + 0.05, 4.85, 0.9,
         [dict(t=[(does, {"size": 15, "color": NAVY}), ("\n" + app, {"size": 13, "color": BLUE, "bold": True})])],
         anchor=MSO_ANCHOR.MIDDLE)
text(s, 10.3, 9.5, 8.76, 0.4,
     [dict(t="Revenue Cloud personas in blue · Field Service personas in dark blue", size=12, color=MUTED)])
notes(s, "Set the scene. Infinitech is an existing demo account. The three products are in the QuantumBit catalog today. Introduce the five personas — the story hands off between them.")

# 3. Demo map
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "The demo in nine scenes")
set_ph(s, 1, "One customer, one asset, one continuous loop · about 15 minutes")
remove_ph(s, 2, 3)

GAP, W, H = 0.38, (18.12 - 4 * 0.38) / 5, 2.45


def map_card(x, y, sc):
    box(s, x, y, W, H)
    circle_num(s, x + 0.2, y + 0.2, sc["n"], d=0.55, fill=BLUE if sc["act"] == 1 else DEEP)
    pill(s, x + W - (1.25 if sc["vision"] else 1.55) - 0.15, y + 0.29,
         "VISION" if sc["vision"] else "CORE DEMO", vision=sc["vision"],
         w=1.25 if sc["vision"] else 1.55, size=11)
    text(s, x + 0.2, y + 0.95, W - 0.4, 0.8, [dict(t=sc["title"], size=19, bold=True, font=HEAD)])
    text(s, x + 0.2, y + 1.7, W - 0.4, 0.65,
         [dict(t=sc["persona"], size=13, color=MUTED)], anchor=MSO_ANCHOR.BOTTOM)


text(s, 0.94, 2.95, 12, 0.4,
     [dict(t="ACT 1 · REVENUE CLOUD → FIELD SERVICE · SELL, DELIVER, BILL", size=14, color=BLUE, bold=True)])
for i, sc in enumerate(SCENES[:5]):
    x = 0.94 + i * (W + GAP)
    map_card(x, 3.4, sc)
    if i < 4:
        arrow(s, x + W + 0.04, 3.4 + H / 2, x + W + GAP - 0.04, 3.4 + H / 2)

text(s, 0.94, 6.2, 12, 0.4,
     [dict(t="ACT 2 · FIELD SERVICE → REVENUE CLOUD · SERVICE, PROTECT, GROW", size=14, color=DEEP, bold=True)])
for i, sc in enumerate(SCENES[5:]):
    x = 0.94 + i * (W + GAP)
    map_card(x, 6.65, sc)
    arrow(s, x + W + 0.04, 6.65 + H / 2, x + W + GAP - 0.04, 6.65 + H / 2, color=DEEP)
x = 0.94 + 4 * (W + GAP)
loop = box(s, x, 6.65, W, H, fill=DEEP)
text(s, x + 0.25, 6.85, W - 0.5, H - 0.4,
     [dict(t="↺  Back to Scene 1", size=19, bold=True, color=WHITE, font=HEAD),
      dict(t="The next quote starts from what the field learned", size=14, color=WHITE, space=8)],
     anchor=MSO_ANCHOR.MIDDLE)

pill(s, 0.94, 9.45, "CORE DEMO", w=1.55, size=11)
text(s, 2.6, 9.43, 5.5, 0.4, [dict(t="Shown live in the org", size=13, color=MUTED)], anchor=MSO_ANCHOR.MIDDLE)
pill(s, 6.2, 9.45, "VISION", vision=True, w=1.25, size=11)
text(s, 7.55, 9.43, 6.5, 0.4, [dict(t="Storyboard only for now · built in a later phase", size=13, color=MUTED)],
     anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Use this as the map. Act 1 flows from Revenue Cloud into Field Service; Act 2 flows back. Scenes 6–8 are vision for now and can be shown as storyboards, or skipped in a live run.")


def section(title, sub):
    sl = prs.slides.add_slide(layout(L_SECTION))
    set_ph(sl, 0, title)
    text(sl, 1.24, 7.75, 17.52, 0.8, [dict(t=sub, size=28, color=WHITE, align=PP_ALIGN.CENTER)])
    return sl


def scene_slide(sc):
    sl = prs.slides.add_slide(layout(L_SCENE))
    text(sl, 0.94, 0.45, 4, 0.4,
         [dict(t=f"SCENE {sc['n']}  ·  ACT {sc['act']}", size=15, color=DEEP if sc["act"] == 2 else BLUE,
               bold=True)])
    set_ph(sl, 0, sc["title"], size=40)
    set_ph(sl, 2, f"{sc['persona']} · {sc['app'].split(' · ')[-1]}", size=20)
    remove_ph(sl, 1, 3)

    text(sl, 0.94, 3.75, 8.3, 0.4, [dict(t="WHAT THE AUDIENCE SEES", size=14, color=BLUE, bold=True)])
    text(sl, 0.94, 4.25, 8.3, 3.4,
         [dict(t=b, size=18, bullet=True, space=0 if i == 0 else 10) for i, b in enumerate(sc["sees"])])

    box(sl, 0.94, 7.9, 8.3, 2.1)
    box(sl, 0.94, 7.9, 0.12, 2.1, fill=ORANGE_LN if sc["vision"] else BLUE, radius=0,
        shape=MSO_SHAPE.RECTANGLE)
    text(sl, 1.35, 8.05, 7.7, 0.4, [dict(t="THE MOMENT", size=13, color=MUTED, bold=True)])
    text(sl, 1.35, 8.45, 7.7, 1.45, [dict(t=sc["moment"], size=20, bold=True, font=HEAD)],
         anchor=MSO_ANCHOR.MIDDLE)

    px = 10.75
    pill(sl, px, 1.05, "VISION" if sc["vision"] else "CORE DEMO", vision=sc["vision"],
         w=1.3 if sc["vision"] else 1.6)
    shot = box(sl, px, 1.7, 8.15, 4.85, fill=PANEL, line=BLUE, dash=True, radius=0.04)
    text(sl, px + 0.4, 1.9, 7.35, 4.45,
         [dict(t="SCREENSHOT PLACEHOLDER", size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER),
          dict(t=sc["shot"], size=17, color=MUTED, align=PP_ALIGN.CENTER, space=10)],
         anchor=MSO_ANCHOR.MIDDLE)

    text(sl, px, 6.85, 8.15, 0.4, [dict(t="CLICK PATH", size=14, color=BLUE, bold=True)])
    for i, step in enumerate(sc["path"]):
        y = 7.35 + i * 0.68
        circle_num(sl, px, y + 0.04, i + 1, d=0.42, fill=ORANGE_LN if sc["vision"] else BLUE, size=12)
        text(sl, px + 0.6, y, 7.5, 0.5, [dict(t=step, size=16)], anchor=MSO_ANCHOR.MIDDLE)

    vision_note = "\n\n[VISION] Not built yet. Present as a storyboard, or skip in a live run." if sc["vision"] else ""
    notes(sl, sc["talk"] + vision_note)
    return sl


section("Act 1 · Revenue Cloud → Field Service",
        "Sell it, deliver it, and start billing when it's actually running")
for sc in SCENES[:5]:
    scene_slide(sc)

section("Act 2 · Field Service → Revenue Cloud",
        "Service it, protect revenue, and grow the account from the field")
for sc in SCENES[5:]:
    scene_slide(sc)

# Why it matters
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "What the audience should walk away with")
set_ph(s, 1, "Three outcomes, each proven by a scene they just watched")
remove_ph(s, 2, 3)
outcomes = [
    ("No billing before go-live", "Billing waits for the technician to complete the install, so invoices match reality.",
     "Scenes 2–5", False),
    ("Service that pays for itself", "Coverage is checked against what was sold. Uncovered work becomes revenue automatically.",
     "Scene 7", True),
    ("Growth that starts in the field", "Technicians flag expansion, and sellers amend the real install base.",
     "Scenes 8–9", True),
]
CW = (18.12 - 2 * 0.45) / 3
for i, (head, body, proof, vis) in enumerate(outcomes):
    x = 0.94 + i * (CW + 0.45)
    box(s, x, 3.2, CW, 6.3)
    circle_num(s, x + 0.4, 3.6, i + 1, d=0.8, size=24, fill=BLUE if i == 0 else DEEP)
    text(s, x + 0.4, 4.75, CW - 0.8, 1.6, [dict(t=head, size=30, bold=True, font=HEAD)])
    text(s, x + 0.4, 6.15, CW - 0.8, 2.2, [dict(t=body, size=19, color=NAVY)])
    text(s, x + 0.4, 8.75, 2.6, 0.4, [dict(t=f"Proven in {proof}", size=14, color=BLUE, bold=True)],
         anchor=MSO_ANCHOR.MIDDLE)
    if vis:
        pill(s, x + CW - 1.65, 8.77, "VISION", vision=True, w=1.25, size=11)
notes(s, "Land the three outcomes. If you skipped the vision scenes, present outcomes 2 and 3 as where the solution goes next.")

# Run of show
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "Run of show")
set_ph(s, 1, "Who drives each scene, how long it runs, and the fallback")
remove_ph(s, 2, 3)
minutes = [2, 1.5, 1.5, 2, 1.5, 1, 2.5, 1.5, 1.5]
fallbacks = [
    "Pre-built quote on Infinitech",
    "Pre-activated order",
    "Pre-created Work Order",
    "Complete from the desktop Work Order",
    "Show a pre-completed plan",
    "Storyboard slide",
    "Storyboard slide",
    "Storyboard slide",
    "Pre-loaded Asset record",
]
rows = len(SCENES) + 2
tbl = s.shapes.add_table(rows, 6, Inches(0.94), Inches(3.05), Inches(18.12), Inches(0.6 * rows)).table
widths = [0.8, 4.4, 3.6, 4.6, 1.1, 3.62]
for i, wd in enumerate(widths):
    tbl.columns[i].width = Inches(wd)
tbl.first_row = True
tbl.horz_banding = False


def cell(r, c, val, size=15, bold=False, color=NAVY, fill=WHITE, align=PP_ALIGN.LEFT):
    ce = tbl.cell(r, c)
    ce.fill.solid()
    ce.fill.fore_color.rgb = fill
    ce.margin_left = ce.margin_right = Inches(0.12)
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = ce.text_frame
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = align
    r_ = p.add_run()
    r_.text = val
    style_run(r_, size, color, bold=bold)


for c, h in enumerate(["#", "Scene", "Persona", "App / screen", "Min", "If it doesn't load"]):
    cell(0, c, h, size=14, bold=True, color=WHITE, fill=NAVY)
for i, sc in enumerate(SCENES, start=1):
    fill = WHITE if i % 2 else LIGHT
    cell(i, 0, str(sc["n"]), bold=True, color=BLUE, fill=fill, align=PP_ALIGN.CENTER)
    cell(i, 1, sc["title"] + ("  (vision)" if sc["vision"] else ""), bold=True,
         color=ORANGE_TX if sc["vision"] else NAVY, fill=fill)
    cell(i, 2, sc["persona"], fill=fill)
    cell(i, 3, sc["app"].split(" · ")[-1] if " · " in sc["app"] else sc["app"], fill=fill)
    cell(i, 4, f"{minutes[i - 1]:g}", fill=fill, align=PP_ALIGN.CENTER)
    cell(i, 5, fallbacks[i - 1], fill=fill, color=MUTED)
last = rows - 1
for c in range(6):
    cell(last, c, "", fill=LIGHT)
cell(last, 1, "Total", bold=True, fill=LIGHT)
cell(last, 4, f"{sum(minutes):g}", bold=True, fill=LIGHT, align=PP_ALIGN.CENTER)
cell(last, 5, f"Core demo only: {sum(m for m, sc in zip(minutes, SCENES) if not sc['vision']):g} min",
     bold=True, color=BLUE, fill=LIGHT)
notes(s, "Timings are targets. Prep the fallbacks in the demo org before the call so any scene can be shown from a pre-built record.")

# Thank you
prs.slides.add_slide(layout(L_THANKS))

prs.save(OUTPUT)
print("saved", OUTPUT, len(prs.slides), "slides")
