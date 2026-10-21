"""Build the Rentokil residential pest control sales/service lifecycle deck
(vision storyboard) from the Salesforce FY27 corporate template."""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor

from deck_helpers import (
    BLUE, BODY, DEEP, HEAD, L_CONTENT, L_SCENE, L_SECTION, L_THANKS, L_TITLE, LIGHT,
    MSO_ANCHOR, MSO_SHAPE, MUTED, NAVY, ORANGE_LN, ORANGE_TX, PANEL, PKG, PKG_BG, PP_ALIGN, WHITE,
    Inches, arrow, box, circle_num, notes, pill, pkg_callout, pkg_pill, remove_ph, set_ph,
    strip_slides, style_run, text,
)

TEMPLATE = sys.argv[1]
OUTPUT = sys.argv[2]

PURPLE = RGBColor(0x5A, 0x1B, 0xA9)

PHASES = [
    dict(key="sell", name="Sell", color=BLUE, tagline="The inspection is the sales call",
         sub="Inspect, price, sign, and book the first treatment in one visit"),
    dict(key="handoff", name="Hand off", color=DEEP, tagline="One order, two kinds of work",
         sub="The signed plan builds the first visit and every visit after it"),
    dict(key="serve", name="Serve", color=NAVY, tagline="Every visit tied to what was sold",
         sub="Billing starts with service, and the guarantee enforces itself"),
    dict(key="grow", name="Grow & retain", color=PURPLE, tagline="The technician is the best salesperson",
         sub="Expand the household plan and keep it through every season"),
]
PHASE = {p["key"]: p for p in PHASES}

SCENES = [
    dict(
        n=1, phase="sell", title="Plan the day",
        persona="Sales Rep", app="Field sales visit plan & map",
        sees=[
            "The rep's morning: inbound leads plotted across Bloomingdale and Glendale Heights",
            "Michael's web request from last night (ants in the kitchen, mice in the garage) is the first stop",
            "Drive time, appointment window, and the lead's notes in one view",
        ],
        moment="Last night's web form is this morning's first stop.",
        look="A map of today's stops with a route line. Tapping 316 Ridgewood Dr opens Michael's lead: "
             "what he reported, when he's home, and the appointment window he picked.",
        feel="Calm and organized. The rep starts the day knowing exactly where to go and why.",
        flow=["Open today's visit plan on the tablet", "Show the route across Bloomingdale",
              "Tap Michael's stop → lead details", "Start the visit on arrival"],
        build="Field sales visit planning and maps (check licensing), lead-to-visit flow",
        talk="""Open on the rep's tablet. Michael filled out a web form last night: ants in the kitchen, mice in the garage. This morning he's the first stop on a route that's already planned.
Point out that everything the rep needs is on the stop: what Michael reported, his window, and the drive time.""",
    ),
    dict(
        n=2, phase="sell", title="Inspect & quote on the tablet",
        persona="Sales Rep", app="Revenue Cloud quote on mobile",
        sees=[
            "The rep walks the property and records findings: ant trails, rodent entry points, a wooded lot",
            "Home details (2,400 sq ft, basement, attached garage) drive attribute-based pricing",
            "Three options appear: Essential, Home Protection, and Home Protection Plus",
            "A rule adds the Initial Service automatically to every recurring plan",
        ],
        moment="The price reflects Michael's home, not a rate card.",
        look="A guided inspection on the tablet: tap pests found, snap photos, enter home size. "
             "Then a side-by-side of three plans, with Home Protection highlighted as recommended.",
        feel="Consultative, not salesy. Michael watches the price build from what was found in his own home.",
        flow=["Capture findings and photos", "Enter home size and property details",
              "Show the three plan options", "Highlight the auto-added Initial Service"],
        build="Pest control catalog, attribute-based pricing, bundle options, mobile-friendly quote capture",
        talk="""The rep walks the house with Michael. Each finding is a tap: ants in the kitchen, droppings near the garage door, a wooded lot behind the house.
Square footage and property details become pricing attributes, so the three options are priced for this home.
Call out the rule: every recurring plan requires an Initial Service, so the rep can't forget it.""",
    ),
    dict(
        n=3, phase="sell", title="Sign & pay at the door",
        persona="Sales Rep + Michael", app="Revenue Cloud order, e-signature, payments",
        sees=[
            "Michael chooses Home Protection and adds seasonal mosquito service (April to October)",
            "He e-signs on the tablet and saves a card for autopay",
            "The order activates on the spot, and a text confirms his first treatment window",
        ],
        moment="Knock to signed in one visit, with the first treatment already booked.",
        look="A clean summary screen: plan, monthly price, initial service price, and start date. "
             "Signature pad, card capture, then a confirmation with the treatment appointment.",
        feel="Fast and trustworthy. No paperwork, no callback, no \"someone from the office will call you.\"",
        flow=["Select Home Protection + mosquito add-on", "Review the summary with Michael",
              "Capture e-signature and card on file", "Activate → confirmation text"],
        build="E-signature, card-on-file payments, mobile order activation, customer notifications",
        talk="""Michael picks Home Protection and adds mosquito service for the summer. He signs on the tablet and saves a card for autopay.
The order activates before the rep leaves the driveway, and Michael gets a text with his first treatment window.
This is the end of the sales cycle and the start of the service relationship, in one visit.""",
    ),
    dict(
        n=4, phase="handoff", title="The plan builds itself",
        persona="Operations", app="Revenue Cloud · Dynamic Revenue Orchestration",
        sees=[
            "The order decomposes into a fulfillment plan automatically",
            "An Initial Service Work Order is created and scheduled within 48 hours",
            "A quarterly maintenance plan and a seasonal mosquito plan are created",
            "Billing waits for the Initial Service to be completed",
        ],
        moment="The rep never calls the office. The service plan builds itself.",
        look="The fulfillment plan as a simple flow: Initial Service, then quarterly visits and mosquito "
             "visits side by side, with Start Billing waiting on the first treatment.",
        feel="Invisible and reliable. Operations watches it happen instead of keying it in.",
        flow=["Open Michael's order → fulfillment plan", "Show the Initial Service Work Order",
              "Show the quarterly and mosquito maintenance plans", "Point at Start Billing waiting"],
        build="Pest DRO overlay (reuses the current Work Order creation flow), subscription → maintenance plan flow",
        pkg="The Initial Service is booked automatically into the best open slot within 48 hours, and Michael "
            "gets his arrival window by text. Nothing waits in a scheduling queue.",
        talk="""Switch to operations. Nobody re-typed the order. Orchestration split it into the first treatment and the recurring plans.
The Initial Service Work Order is already in the queue for the next 48 hours, and the quarterly and mosquito plans exist from day one.
Billing is deliberately waiting on the first treatment.""",
    ),
    dict(
        n=5, phase="serve", title="First treatment starts billing",
        persona="Technician", app="Field Service Mobile",
        sees=[
            "The technician arrives with the rep's inspection notes and photos on the job",
            "Treats the kitchen, seals two garage entry points, and logs the products applied",
            "Completing the job starts the subscription and charges the Initial Service to Michael's card",
        ],
        moment="Michael pays when service starts, not when he signs.",
        look="The mobile job card: address, Michael's findings and photos, a checklist of treatment areas, "
             "products used, and a Complete button that triggers a service summary to Michael.",
        feel="Prepared and professional. The technician already knows the house before ringing the bell.",
        flow=["Open the job on mobile", "Review the rep's photos and findings",
              "Log treatment areas and products used", "Complete → service summary + first charge"],
        build="Reuses the current completion flow; adds the billing trigger and the customer service summary",
        pkg="Appointment Assistant texts Michael when the technician is on the way, with live location and ETA. "
            "No more waiting at home all afternoon.",
        talk="""The technician arrives already knowing the house: the rep's photos and notes are on the job.
They treat the kitchen, seal the garage, and log what was applied, which also feeds compliance records.
Completing the job closes the fulfillment step, starts the subscription, and charges the Initial Service.""",
    ),
    dict(
        n=6, phase="serve", title="Recurring visits run themselves",
        persona="Dispatcher", app="Field Service scheduling",
        sees=[
            "Quarterly visits are generated from the maintenance plan without anyone booking them",
            "Each visit is routed into the Bloomingdale territory alongside neighboring customers",
            "Michael gets a reminder text with an arrival window two days before",
        ],
        moment="Recurring revenue and recurring routes come from the same record.",
        look="The dispatcher's schedule for a Bloomingdale route: a dense cluster of short residential visits, "
             "with Michael's quarterly visit slotted in between neighbors.",
        feel="Efficient. Density turns into margin, and Michael never has to call to schedule.",
        flow=["Open the maintenance plan → upcoming visits", "Show the generated quarterly Work Order",
              "Show it on the Bloomingdale route", "Show the reminder text to Michael"],
        build="Maintenance plan generation, territory routing, appointment reminders",
        pkg="Optimization packs every quarterly visit into dense Bloomingdale routes and cuts drive time. "
            "This is where residential margin comes from, and Core alone can't show it.",
        talk="""Every quarter, the maintenance plan generates Michael's visit on its own. Dispatch sees it in the Bloomingdale route next to his neighbors.
That density is where residential pest control makes money. Michael just gets a reminder text.""",
    ),
    dict(
        n=7, phase="serve", title="Re-service with a guarantee check",
        persona="Customer Care + Technician", app="Field Service + Revenue Cloud Billing",
        sees=[
            "Three weeks later, Michael reports ants are back in the kitchen",
            "A guarantee check finds his active Home Protection plan: re-service at no charge",
            "A neighbor without a plan gets a priced, one-time treatment that is invoiced automatically",
        ],
        moment="The guarantee is enforced by the subscription, not by memory.",
        look="A re-service request with a clear banner: \"Covered by Home Protection, no charge.\" "
             "Beside it, the neighbor's request showing a one-time treatment price and invoice.",
        feel="Fair and confident. Care agents never have to guess what a customer is entitled to.",
        flow=["Log Michael's re-service request", "Show the Covered result and $0 Work Order",
              "Repeat for the neighbor without a plan", "Show the one-time order and invoice"],
        build="Guarantee check flow keyed on an active subscription, one-time treatment product and price",
        pkg="Appointment booking offers Michael real open windows while he's still on the phone, and slots the "
            "re-service into a technician route already nearby.",
        talk="""Ants come back three weeks later. That's normal, and it's what the guarantee is for.
The check looks at Michael's active subscription and creates a no-charge re-service. A neighbor without a plan gets a priced one-time treatment and an invoice.
Same process, the right outcome for each customer, decided by what they bought.""",
    ),
    dict(
        n=8, phase="grow", title="Upsell in the driveway",
        persona="Technician → Michael", app="Field Service Mobile + Revenue Cloud amendment",
        sees=[
            "On a quarterly visit, the technician spots mud tubes on the foundation: termite activity",
            "They photograph it and send Michael a termite plan quote from the job",
            "Michael approves on his phone; the amendment co-terms with his existing plan",
            "His next visit includes the termite inspection automatically",
        ],
        moment="The person at the house is the best salesperson the company has.",
        look="On the job card, a \"Recommend service\" action with the photo attached. Michael receives a "
             "short quote by text and approves it with one tap.",
        feel="Helpful, not pushy. The recommendation comes with evidence from Michael's own foundation.",
        flow=["Quarterly visit → Recommend service", "Attach the termite photo and pick Termite Plan",
              "Michael approves by text", "Amendment co-terms → next visit adds termite inspection"],
        build="Recommend-service quick action, mobile-approved quote, subscription amendment",
        talk="""On a routine quarterly visit, the technician finds mud tubes on the foundation. They take a photo and send Michael a termite plan from the job.
Michael approves from his phone. Revenue Cloud amends the existing subscription, so it co-terms and bills together, and the next visit includes the termite inspection.""",
    ),
    dict(
        n=9, phase="grow", title="Pause, resume & renew",
        persona="Billing + Michael", app="Revenue Cloud subscription lifecycle",
        sees=[
            "Mosquito service pauses in November; visits and billing stop together",
            "It resumes automatically in April with no call and no re-sign",
            "At the one-year mark, the plan renews with the annual price adjustment",
            "One subscription view shows every visit, invoice, and change",
        ],
        moment="Seasonal service without seasonal churn.",
        look="Michael's subscription timeline: Home Protection running all year, mosquito shaded April to "
             "October each year, termite added in month five, and a renewal marker at month twelve.",
        feel="Effortless for Michael, predictable for finance.",
        flow=["Open Michael's subscription timeline", "Show the mosquito pause and resume",
              "Show the termite amendment", "Show the renewal quote and price adjustment"],
        build="Seasonal ramp or suspend/resume, renewal with price uplift, subscription timeline view",
        talk="""Close on the subscription. Mosquito service pauses for the winter and comes back in April without a call.
At one year, the whole plan renews with the annual adjustment. Everything we showed, from the inspection to the termite add-on, lives on this one record.""",
    ),
]

prs = Presentation(TEMPLATE)
strip_slides(prs)
M = prs.slide_masters[1]


def layout(i):
    return M.slide_layouts[i]


# 1. Title
s = prs.slides.add_slide(layout(L_TITLE))
set_ph(s, 0, "One subscription, every visit")
set_ph(s, 1, "Revenue Cloud + Field Service: the residential pest control lifecycle, with Rentokil as the example")
set_ph(s, 2, "Vision storyboard · Draft")
notes(s, "This deck shows how a residential pest control company like Rentokil would run the full sales and service lifecycle on Revenue Cloud and Field Service. Everything here is vision: it describes how the demo would look and feel, not what's built today.")

# 2. Meet Michael
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "Meet Michael")
set_ph(s, 1, "Ants in the kitchen, mice in the garage, and a wooded lot")
remove_ph(s, 2, 3)

box(s, 0.94, 3.1, 8.6, 2.75)
circle_num(s, 1.25, 3.4, "M", d=0.9, fill=BLUE, size=26)
text(s, 2.4, 3.35, 6.9, 0.55, [dict(t="Michael", size=26, bold=True, font=HEAD)])
text(s, 2.4, 3.95, 6.9, 0.9,
     [dict(t="316 Ridgewood Dr", size=17), dict(t="Bloomingdale, IL 60108", size=17)])
text(s, 1.25, 4.9, 8.0, 0.85,
     [dict(t="Single-family home · about 2,400 sq ft · basement and attached garage · wooded lot",
           size=15, color=MUTED)])

text(s, 0.94, 6.15, 8.6, 0.4, [dict(t="WHAT HE ENDS UP WITH", size=14, color=BLUE, bold=True)])
plan = [
    ("Home Protection plan", "Quarterly general pest and rodent service", "Scene 3"),
    ("Initial Service", "One-time cleanout and exclusion, added by rule", "Scene 2"),
    ("Seasonal mosquito", "April to October, pauses every winter", "Scenes 3, 9"),
    ("Termite plan", "Added in the driveway on a quarterly visit", "Scene 8"),
]
for i, (name, desc, where) in enumerate(plan):
    y = 6.6 + i * 0.86
    box(s, 0.94, y, 8.6, 0.74)
    text(s, 1.2, y + 0.04, 3.4, 0.66, [dict(t=name, size=16, bold=True)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.55, y + 0.04, 3.65, 0.66, [dict(t=desc, size=13, color=MUTED)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.1, y + 0.04, 1.3, 0.66, [dict(t=where, size=12, color=BLUE, bold=True, align=PP_ALIGN.RIGHT)],
         anchor=MSO_ANCHOR.MIDDLE)

text(s, 10.3, 3.05, 8.7, 0.4, [dict(t="THE CAST", size=14, color=BLUE, bold=True)])
cast = [
    ("Sales Rep", "Inspects, quotes, and signs at the door", "Field sales + Revenue Cloud", BLUE),
    ("Operations", "Watches the plan build itself", "Revenue Cloud orchestration", DEEP),
    ("Dispatcher", "Routes visits by territory", "Field Service", NAVY),
    ("Technician", "Treats, re-services, spots upsells", "Field Service Mobile", NAVY),
    ("Customer Care & Billing", "Guarantees, pauses, renewals", "Revenue Cloud Billing", DEEP),
    ("Michael", "Signs, approves, and gets texts", "Customer experience", PURPLE),
]
for i, (role, does, app, col) in enumerate(cast):
    y = 3.55 + i * 1.0
    box(s, 10.3, y, 8.76, 0.86)
    circle_num(s, 10.5, y + 0.15, role[0], d=0.56, fill=col, size=16)
    text(s, 11.3, y + 0.04, 3.9, 0.78, [dict(t=role, size=17, bold=True)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 14.4, y + 0.04, 4.55, 0.78,
         [dict(t=[(does, {"size": 14, "color": NAVY}), ("\n" + app, {"size": 12, "color": BLUE, "bold": True})])],
         anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Michael is the customer for the whole story. He starts with a simple pest problem and, over a year, becomes a multi-service household. The plan names and prices are illustrative.")

# 3. Lifecycle at a glance
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "The lifecycle at a glance")
set_ph(s, 1, "Four phases, nine scenes, one household · about 16 minutes")
remove_ph(s, 2, 3)

GAP = 0.4
CW = (18.12 - 3 * GAP) / 4
for i, ph in enumerate(PHASES):
    x = 0.94 + i * (CW + GAP)
    box(s, x, 3.05, CW, 6.15)
    head = box(s, x, 3.05, CW, 1.35, fill=ph["color"], radius=0.12)
    text(s, x + 0.3, 3.15, CW - 0.6, 1.15,
         [dict(t=ph["name"], size=24, bold=True, color=WHITE, font=HEAD),
          dict(t=ph["tagline"], size=13, color=WHITE, space=2)],
         anchor=MSO_ANCHOR.MIDDLE)
    scenes = [sc for sc in SCENES if sc["phase"] == ph["key"]]
    for j, sc in enumerate(scenes):
        y = 4.7 + j * 1.45
        circle_num(s, x + 0.3, y, sc["n"], d=0.5, fill=ph["color"], size=14)
        persona_runs = [(sc["persona"], {})]
        if sc.get("pkg"):
            persona_runs.append(("  ✦ package", {"color": PKG, "bold": True}))
        text(s, x + 0.95, y - 0.04, CW - 1.2, 1.3,
             [dict(t=sc["title"], size=17, bold=True),
              dict(t=persona_runs, size=13, color=MUTED, space=2)])
    if ph["key"] == "handoff":
        text(s, x + 0.3, 6.1, CW - 0.6, 0.4, [dict(t="ONE ORDER CREATES", size=12, color=MUTED, bold=True)])
        text(s, x + 0.3, 6.5, CW - 0.6, 2.6,
             [dict(t=b, size=14, bullet=True, space=0 if k == 0 else 6) for k, b in enumerate([
                 "Initial Service Work Order within 48 hours",
                 "Quarterly maintenance plan",
                 "Seasonal mosquito plan",
                 "Billing that waits for the first treatment",
             ])])
    if i < 3:
        arrow(s, x + CW + 0.05, 3.72, x + CW + GAP - 0.05, 3.72, color=ph["color"])

loop = box(s, 0.94, 9.45, 18.12, 0.6, fill=LIGHT, radius=0.5)
text(s, 1.3, 9.45, 9.5, 0.6,
     [dict(t="↺  Renewal and next season's service start the loop again", size=15, bold=True, color=DEEP)],
     anchor=MSO_ANCHOR.MIDDLE)
text(s, 10.6, 9.45, 4.9, 0.6,
     [dict(t=[("✦ package", {"color": PKG, "bold": True}),
              ("  enhanced by the FSL managed package", {"color": MUTED})], size=13)],
     anchor=MSO_ANCHOR.MIDDLE)
pill(s, 15.7, 9.57, "ALL SCENES: VISION", vision=True, w=2.3, size=11)
notes(s, "Use this as the map. Sell happens at Michael's door. Hand off is invisible automation. Serve is every visit after that. Grow and retain is where the technician and the subscription keep and expand the household. Scenes 4–7 are where the Field Service managed package adds the most: automatic booking, live arrival updates, route optimization, and appointment booking.")

# 4. Field Service Core vs managed package
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "Field Service Core vs the managed package")
set_ph(s, 1, "Core runs the lifecycle; the package makes residential routes pay")
remove_ph(s, 2, 3)
cols = [
    ("FIELD SERVICE CORE", "Included with Field Service licenses", BLUE, [
        "Work Orders, Service Appointments, and Assets",
        "Maintenance Plans that generate recurring visits",
        "Skills, territories, operating hours, Work Types",
        "Entitlements for guarantee checks",
        "Field Service mobile app for technicians",
        "Scheduling is manual: set the time, add the technician",
    ]),
    ("MANAGED PACKAGE ADDS", "Field Service app from AppExchange (FSL)", PKG, [
        "Dispatcher Console: Gantt, map, drag and drop",
        "Automatic booking and ranked technician candidates",
        "Rules that enforce skills, territory, hours, and travel",
        "Route optimization across the whole territory",
        "Appointment booking with real arrival windows",
        "Appointment Assistant: live \"on my way\" updates",
    ]),
    ("WHERE IT SHOWS UP", "Scenes that change with the package", DEEP, [
        "Scene 4: Initial Service booked within 48 hours, automatically",
        "Scene 5: Michael tracks the technician's arrival live",
        "Scene 6: quarterly visits optimized into dense routes",
        "Scene 7: re-service booked into a real open window",
        "Scenes 1–3, 8–9: unchanged; the Revenue Cloud story is all Core",
    ]),
]
C3 = (18.12 - 2 * 0.4) / 3
for i, (head, sub, col, items) in enumerate(cols):
    x = 0.94 + i * (C3 + 0.4)
    box(s, x, 3.05, C3, 6.05)
    box(s, x, 3.05, C3, 0.12, fill=col, radius=0, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.35, 3.35, C3 - 0.7, 0.4, [dict(t=head, size=15, color=col, bold=True)])
    text(s, x + 0.35, 3.78, C3 - 0.7, 0.4, [dict(t=sub, size=13, color=MUTED)])
    text(s, x + 0.35, 4.4, C3 - 0.7, 4.6,
         [dict(t=b, size=16, bullet=True, space=0 if k == 0 else 9) for k, b in enumerate(items)])
box(s, 0.94, 9.35, 18.12, 0.75, fill=PKG_BG, line=PKG, radius=0.3, line_w=1)
text(s, 1.3, 9.35, 17.4, 0.75,
     [dict(t=[("For residential pest control, the package is close to essential. ", {"bold": True, "color": PKG}),
              ("Route density is the margin story (Scene 6), and Core can't show it. Plan for post-install setup: "
               "package permission sets and a default scheduling policy.", {"color": NAVY})], size=14)],
     anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Be clear about what's Core and what's the package. Revenue Cloud plus Core Field Service runs the whole lifecycle. The managed package is what makes dispatch, routing, and arrival updates look the way a pest control operations leader expects. If the package isn't installed in the demo org, present Scenes 4–7 with the callouts as what the package adds.")


def section(ph):
    sl = prs.slides.add_slide(layout(L_SECTION))
    n = PHASES.index(ph) + 1
    text(sl, 1.24, 3.35, 17.52, 0.6,
         [dict(t=f"PHASE {n}  ·  {ph['name'].upper()}", size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)])
    set_ph(sl, 0, ph["tagline"], size=60)
    text(sl, 1.24, 7.75, 17.52, 0.8, [dict(t=ph["sub"], size=28, color=WHITE, align=PP_ALIGN.CENTER)])
    notes(sl, ph["sub"])


def scene_slide(sc):
    ph = PHASE[sc["phase"]]
    sl = prs.slides.add_slide(layout(L_SCENE))
    text(sl, 0.94, 0.45, 6, 0.4,
         [dict(t=f"SCENE {sc['n']}  ·  {ph['name'].upper()}", size=15, color=ph["color"], bold=True)])
    set_ph(sl, 0, sc["title"], size=40)
    set_ph(sl, 2, f"{sc['persona']} · {sc['app'].split(' · ')[-1]}", size=20)
    remove_ph(sl, 1, 3)

    text(sl, 0.94, 3.75, 8.3, 0.4, [dict(t="WHAT THE AUDIENCE SEES", size=14, color=BLUE, bold=True)])
    text(sl, 0.94, 4.25, 8.3, 3.5,
         [dict(t=b, size=17, bullet=True, space=0 if i == 0 else 9) for i, b in enumerate(sc["sees"])])

    box(sl, 0.94, 7.9, 8.3, 2.1)
    box(sl, 0.94, 7.9, 0.12, 2.1, fill=ph["color"], radius=0, shape=MSO_SHAPE.RECTANGLE)
    text(sl, 1.35, 8.05, 7.7, 0.4, [dict(t="THE MOMENT", size=13, color=MUTED, bold=True)])
    text(sl, 1.35, 8.45, 7.7, 1.45, [dict(t=sc["moment"], size=20, bold=True, font=HEAD)],
         anchor=MSO_ANCHOR.MIDDLE)

    px = 10.75
    pill(sl, px, 1.05, "VISION", vision=True, w=1.3)
    pkg = sc.get("pkg")
    look_h = 3.35 if pkg else 4.85
    if pkg:
        pkg_pill(sl, px + 1.45, 1.05)
    box(sl, px, 1.7, 8.15, look_h, fill=PANEL, line=BLUE, dash=True, radius=0.04)
    text(sl, px + 0.4, 1.85, 7.35, look_h - 0.3,
         [dict(t="WHAT IT WOULD LOOK LIKE", size=13, color=BLUE, bold=True),
          dict(t=sc["look"], size=15 if pkg else 16, color=NAVY, space=6 if pkg else 8),
          dict(t=[("Feels like: ", {"bold": True, "color": ORANGE_TX}), (sc["feel"], {"color": MUTED})],
               size=14 if pkg else 15, space=10 if pkg else 14)],
         anchor=MSO_ANCHOR.MIDDLE)
    if pkg:
        pkg_callout(sl, px, 5.2, 8.15, 1.45, pkg, size=14)

    text(sl, px, 6.85, 8.15, 0.4, [dict(t="HOW IT FLOWS", size=14, color=BLUE, bold=True)])
    for i, step in enumerate(sc["flow"]):
        y = 7.35 + i * 0.68
        circle_num(sl, px, y + 0.04, i + 1, d=0.42, fill=ph["color"], size=12)
        text(sl, px + 0.6, y, 7.5, 0.5, [dict(t=step, size=16)], anchor=MSO_ANCHOR.MIDDLE)

    pkg_note = f"\n\n[MANAGED PACKAGE] {pkg}" if pkg else ""
    notes(sl, sc["talk"] + pkg_note + f"\n\n[VISION] To make this live: {sc['build']}.")


for ph in PHASES:
    section(ph)
    for sc in SCENES:
        if sc["phase"] == ph["key"]:
            scene_slide(sc)

# Outcomes
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "What the audience should walk away with")
set_ph(s, 1, "Three outcomes, one for each stage of the household relationship")
remove_ph(s, 2, 3)
outcomes = [
    ("Close on the first visit", "Inspect, price, sign, pay, and book the first treatment before leaving the driveway.",
     "Scenes 1–3", BLUE),
    ("Every visit is accountable", "Billing starts with service, visits generate themselves, and the guarantee is enforced by the plan.",
     "Scenes 4–7", NAVY),
    ("Grow and keep every household", "Technicians add services with evidence, and seasonal plans pause instead of churning.",
     "Scenes 8–9", PURPLE),
]
CW3 = (18.12 - 2 * 0.45) / 3
for i, (head, body, proof, col) in enumerate(outcomes):
    x = 0.94 + i * (CW3 + 0.45)
    box(s, x, 3.2, CW3, 6.3)
    circle_num(s, x + 0.4, 3.6, i + 1, d=0.8, size=24, fill=col)
    text(s, x + 0.4, 4.75, CW3 - 0.8, 1.6, [dict(t=head, size=30, bold=True, font=HEAD)])
    text(s, x + 0.4, 6.15, CW3 - 0.8, 2.2, [dict(t=body, size=19, color=NAVY)])
    text(s, x + 0.4, 8.75, 2.6, 0.4, [dict(t=f"Shown in {proof}", size=14, color=col, bold=True)],
         anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Land the three outcomes in the customer's language: close faster, run every visit profitably, and keep households longer.")

# Run of show / path to live
s = prs.slides.add_slide(layout(L_CONTENT))
set_ph(s, 0, "Run of show and path to live")
set_ph(s, 1, "Who drives each scene, timing, and what it needs to go live")
remove_ph(s, 2, 3)
minutes = [1.5, 2.5, 1.5, 1.5, 2, 1.5, 2, 2, 1.5]
rows = len(SCENES) + 2
tbl = s.shapes.add_table(rows, 6, Inches(0.94), Inches(3.05), Inches(18.12), Inches(0.6 * rows)).table
for i, wd in enumerate([0.7, 4.0, 3.3, 0.9, 1.5, 7.72]):
    tbl.columns[i].width = Inches(wd)
tbl.first_row = True
tbl.horz_banding = False


def cell(r, c, val, size=14, bold=False, color=NAVY, fill=WHITE, align=PP_ALIGN.LEFT):
    ce = tbl.cell(r, c)
    ce.fill.solid()
    ce.fill.fore_color.rgb = fill
    ce.margin_left = ce.margin_right = Inches(0.12)
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = ce.text_frame
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = align
    r_ = p.add_run()
    r_.text = val
    style_run(r_, size, color, bold=bold)


for c, h in enumerate(["#", "Scene", "Persona", "Min", "Package", "To make it live"]):
    cell(0, c, h, bold=True, color=WHITE, fill=NAVY)
for i, sc in enumerate(SCENES, start=1):
    fill = WHITE if i % 2 else LIGHT
    cell(i, 0, str(sc["n"]), bold=True, color=PHASE[sc["phase"]]["color"], fill=fill, align=PP_ALIGN.CENTER)
    cell(i, 1, sc["title"], bold=True, fill=fill)
    cell(i, 2, sc["persona"], fill=fill)
    cell(i, 3, f"{minutes[i - 1]:g}", fill=fill, align=PP_ALIGN.CENTER)
    cell(i, 4, "✦ Enhances" if sc.get("pkg") else "", bold=True, color=PKG, fill=fill, size=12)
    cell(i, 5, sc["build"], size=12, color=MUTED, fill=fill)
last = rows - 1
for c in range(6):
    cell(last, c, "", fill=LIGHT)
cell(last, 1, "Total", bold=True, fill=LIGHT)
cell(last, 3, f"{sum(minutes):g}", bold=True, fill=LIGHT, align=PP_ALIGN.CENTER)
cell(last, 4, "Scenes 4–7", bold=True, color=PKG, fill=LIGHT, size=12)
cell(last, 5, "Scenes 4 and 5 reuse the Work Order creation and completion flows already built",
     bold=True, color=BLUE, fill=LIGHT, size=13)
notes(s, "Every scene is vision today. The right-hand column is the build list. The orchestration-to-Work Order handoff and the Work Order completion flow from the QuantumBit build carry over directly.")

prs.slides.add_slide(layout(L_THANKS))

prs.save(OUTPUT)
print("saved", OUTPUT, len(prs.slides), "slides")
