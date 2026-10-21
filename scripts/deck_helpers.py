"""Shared drawing helpers for the Salesforce-template demo story decks."""
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x03, 0x2D, 0x60)
BLUE = RGBColor(0x06, 0x6A, 0xFE)
DEEP = RGBColor(0x02, 0x2A, 0xC0)
LIGHT = RGBColor(0xEA, 0xF5, 0xFE)
PANEL = RGBColor(0xF3, 0xF8, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x6B, 0x85)
ORANGE_BG = RGBColor(0xFF, 0xEB, 0xD6)
ORANGE_TX = RGBColor(0xA8, 0x4B, 0x00)
ORANGE_LN = RGBColor(0xF3, 0x83, 0x04)

HEAD = "Avant Garde Demi SFDC"
BODY = "Salesforce Sans"

L_TITLE, L_CONTENT, L_SCENE, L_SECTION, L_THANKS = 0, 4, 16, 35, 56

# ---------------------------------------------------------------- helpers


def strip_slides(prs):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)


def remove_ph(slide, *idxs):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx in idxs:
            ph._element.getparent().remove(ph._element)


def set_ph(slide, idx, text, size=None):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text_frame.text = text
            for para in ph.text_frame.paragraphs:
                pPr = para._p.get_or_add_pPr()
                pPr.set("marL", "0")
                pPr.set("indent", "0")
                if size:
                    for r in para.runs:
                        r.font.size = Pt(size)
            return ph
    raise KeyError(idx)


def style_run(run, size, color, bold=False, font=BODY, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def bullet(paragraph, char="•", indent_in=0.3):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT,
         margin=0.0):
    """paras: list of dicts {t, size, color, bold, font, bullet, space}"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        if p.get("space"):
            para.space_before = Pt(p["space"])
        runs = p["t"] if isinstance(p["t"], list) else [(p["t"], {})]
        for chunk, over in runs:
            r = para.add_run()
            r.text = chunk
            style_run(r, over.get("size", p.get("size", 18)),
                      over.get("color", p.get("color", NAVY)),
                      over.get("bold", p.get("bold", False)),
                      over.get("font", p.get("font", BODY)),
                      over.get("italic", p.get("italic", False)))
        if p.get("bullet"):
            bullet(para, p.get("char", "•"))
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        dash=False, line_w=1.25):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4  # dash
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def pill(slide, x, y, label, vision=False, w=None, size=12):
    w = w or (0.22 + 0.105 * len(label))
    fill, color, line = (ORANGE_BG, ORANGE_TX, ORANGE_LN) if vision else (BLUE, WHITE, None)
    s = box(slide, x, y, w, 0.36, fill=fill, line=line, radius=0.5, line_w=1)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    style_run(r, size, color, bold=True)
    return s


def circle_num(slide, x, y, n, d=0.55, fill=BLUE, size=16):
    s = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    style_run(r, size, WHITE, bold=True, font=HEAD)
    return s


def arrow(slide, x1, y1, x2, y2, color=BLUE, w=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()
