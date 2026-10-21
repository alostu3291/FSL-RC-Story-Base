"""Fill the demo story deck with screenshots from screenshots/ and align its text with the built demo.

Rebuilds deck/<name>.pptx from the unfilled storyboard in deck/source, so it can be rerun safely:
    python3 scripts/fill_deck.py
"""
import copy
import json
import math
import os

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "deck", "QuantumBit Revenue Cloud + Field Service Demo Story.pptx")
SRC = os.path.join(ROOT, "deck", "source", "QuantumBit Revenue Cloud + Field Service Demo Story - Storyboard.pptx")
SHOTS = os.path.join(ROOT, "screenshots")
NAVY = RGBColor(3, 45, 96)
GREY = RGBColor(90, 107, 133)
BORDER = RGBColor(201, 201, 201)

# scene number -> (slide number, hero screenshot, top offset px for tall shots)
HERO = {
    1: (5, "s1-02-quote", None),
    2: (6, "s2-02-fulfillment-plan", None),
    3: (7, "s3-01-install-work-order", None),
    4: (8, "s4-03-asset-serial-number", None),
    5: (9, "s5-03-order-fulfillment-complete", None),
    6: (11, "s6-02-maintenance-work-orders", 0),
    7: (12, "s7-03-uncovered-completed-billable-order", None),
    8: (13, "s8-01-flag-expansion-action", None),
    9: (14, "s9-02-asset-related", 1280),
}

# Tall full-page shots in the appendix: top of the band to show, in screenshot pixels.
APPENDIX_CROP = {
    "s3-03-order-work-orders": 720,
    "s4-02-appointment-assigned-resource": 900,
    "s9-02-asset-related": 1500,
    "s9-03-account-related": 620,
}

TITLES = {1: "Configure & quote", 2: "Activate the order", 3: "Automatic dispatch", 4: "Technician on site",
          5: "Billing starts at install", 6: "Maintenance runs itself", 7: "Break/fix with coverage check",
          8: "Upsell from the field", 9: "One install base"}

CAPTIONS = {
    "s1-01-opportunity": "Infinitech opportunity",
    "s1-02-quote": "Quote with hardware, install and Gold",
    "s1-03-quote-lines": "Quote lines and pricing",
    "s1-04-order": "Activated order",
    "s2-01-order-fulfillment-ship-pending": "Order Fulfillment tab: plan waiting on Ship Hardware",
    "s2-02-fulfillment-plan": "Fulfillment plan",
    "s3-01-install-work-order": "Install Work Order linked to Order, step and Asset",
    "s3-02-service-appointment": "Service Appointment in the San Francisco territory",
    "s3-03-order-work-orders": "Installation Work Orders on the Order",
    "s4-01-appointment-scheduled": "Appointment scheduled",
    "s4-02-appointment-assigned-resource": "Technician assigned",
    "s4-03-asset-serial-number": "Serial number on the Asset",
    "s5-01-billing-waiting-for-install": "Billing schedule waiting for the install milestone",
    "s5-02-work-order-completed": "Install Work Order completed",
    "s5-03-order-fulfillment-complete": "Plan after the install: every step completed",
    "s5-04-billing-ready": "R750 billing schedule: Completely Billed",
    "s5-05-invoice-posted": "Invoice posted automatically",
    "s5-06-invoice-lines": "Invoice lines: hardware, install, Gold",
    "s6-01-maintenance-plan": "Gold Maintenance Plan on the Asset",
    "s6-02-maintenance-work-orders": "Generated preventive Work Orders",
    "s6-03-preventive-work-order": "Preventive maintenance Work Order",
    "s7-01-break-fix-covered": "Break/fix on the Gold asset: Covered",
    "s7-02-break-fix-not-covered": "Break/fix for Apex Dynamics, no Gold: Not Covered",
    "s7-03-uncovered-completed-billable-order": "Completed uncovered repair with its Billable Order",
    "s7-04-repair-order": "Repair order",
    "s7-05-repair-invoice": "Repair invoice posted",
    "s8-01-flag-expansion-action": "Flag Expansion from the Work Order",
    "s8-02-flag-expansion-confirmation": "Opportunity created",
    "s8-03-expansion-opportunity": "Expansion opportunity for the seller",
    "s8-04-account-assets-amend": "Account Assets tab: amend the install base",
    "s9-01-asset-details": "R750 Asset details",
    "s9-02-asset-related": "Asset: Work Orders, maintenance, entitlements",
    "s9-03-account-related": "Account: one order, its opportunities and billing",
}

# slide number -> shape id -> list of paragraph texts
TEXT = {
    5: {8: ["A seller builds the Infinitech quote: PowerSwerve R750 server bundle",
            "Installation and Assembly (QB-INST-1000) is quoted with the hardware as a one-time service",
            "Gold Hardware Maintenance adds a year of monthly visits and break/fix coverage"],
        22: ["Add Installation and Assembly"],
        24: ["Add Gold Maintenance → Create Order"]},
    9: {8: ["Completing the Work Order completes the fulfillment step automatically",
            "Installation Complete accomplishes the billing milestone on the hardware and install lines",
            "The invoice posts on its own, the moment the install is done"],
        18: ["Complete the appointment, then the Work Order"],
        20: ["Fulfillment tab → Installation Complete is done"],
        22: ["Billing schedule moves out of Waiting for Milestone"],
        24: ["Open the posted invoice"]},
    11: {18: ["Open the Maintenance Plan from the R750 Asset"],
         22: ["Generate Work Orders → monthly preventive visits"]},
    12: {20: ["Repeat for Apex Dynamics (no Gold maintenance) → Not Covered"],
         22: ["Complete it → a repair order fills Billable Order"],
         24: ["Show the posted repair invoice"]},
    13: {20: ["Opportunity appears for the account team"],
         22: ["Seller amends from the Account Assets tab"]},
}


def set_paragraphs(shape, lines):
    paras = shape.text_frame.paragraphs
    for i, text in enumerate(lines):
        if i >= len(paras):
            new = copy.deepcopy(paras[-1]._p)
            paras[-1]._p.addnext(new)
            paras = shape.text_frame.paragraphs
        runs = paras[i].runs
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)


def by_id(slide, shape_id):
    return next(s for s in slide.shapes if s.shape_id == shape_id)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def content_height(path):
    """Height of the page content, ignoring the empty background below it on tall shots."""
    img = Image.open(path).convert("RGB")
    w, h = img.size
    bg = img.getpixel((w // 2, h - 1))
    diff = ImageChops.difference(img, Image.new("RGB", img.size, bg)).convert("L").point(lambda v: 255 if v > 12 else 0)
    box = diff.getbbox()
    return min(h, box[3] + 20) if box else h


def add_picture(slide, name, left, top, width, height, top_px=None):
    """Fit the whole screenshot in the box, or crop a box-shaped band starting at top_px."""
    path = os.path.join(SHOTS, name + ".png")
    w, h = Image.open(path).size
    if top_px is None:
        start, band = 0, content_height(path)
    else:
        start, band = top_px, min(h - top_px, w * height / width)
    scale = min(width / w, height / band)
    pic_w, pic_h = int(w * scale), int(band * scale)
    pic = slide.shapes.add_picture(path, int(left + (width - pic_w) / 2), int(top), pic_w, pic_h)
    pic.crop_top = start / h
    pic.crop_bottom = (h - start - band) / h
    pic.line.color.rgb = BORDER
    pic.line.width = Pt(0.75)
    return pic


def badge_core(slide, shape_id, width):
    shape = by_id(slide, shape_id)
    set_paragraphs(shape, ["CORE DEMO"])
    right = shape.left + shape.width
    shape.width = Inches(width)
    shape.left = right - shape.width


BLUE = RGBColor(0x06, 0x6A, 0xFE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE_FILLS = {"F38304", "FFEBD6"}
VISION_NOTE = "\n\n[VISION] Not built yet. Present as a storyboard, or skip in a live run."
NOTES = {
    1: ("Purpose: walk the audience through one customer lifecycle that crosses Revenue Cloud and Field Service "
        "in both directions. Every slide after this maps to something we show live.\n"
        "Story: high tech, data center. QuantumBit sells PowerSwerve rack servers to Infinitech for its "
        "San Francisco data center.\n"
        "The sale happens the typical way: a seller, over the phone or as the account rep, quotes and orders in "
        "Revenue Cloud. Technicians never sell in this story; Field Service starts once the order is activated."),
    5: ("Start with the seller, working the deal over the phone or as Infinitech's account rep, the way this "
        "customer normally buys. Infinitech is expanding its data center and wants QuantumBit servers.\n"
        "Add the PowerSwerve R750 bundle, then Installation and Assembly as a one-time service so the install "
        "is sold with the hardware.\n"
        "Add Gold Hardware Maintenance at its annual price. Every line we sell here becomes work for Field "
        "Service later."),
}
NOTE_EDITS = {
    3: ("Scenes 6–8 are vision for now and can be shown as storyboards, or skipped in a live run.",
        "All nine scenes run live in the org."),
    15: (" If you skipped the vision scenes, present outcomes 2 and 3 as where the solution goes next.", ""),
}


def recolor_vision(slide):
    """Vision styling (orange) becomes the core demo styling (blue) now that scenes 6-8 are built."""
    for sh in slide.shapes:
        try:
            fill = str(sh.fill.fore_color.rgb) if sh.fill.type == 1 else None
        except (AttributeError, TypeError):
            fill = None
        if fill in ORANGE_FILLS:
            sh.fill.solid()
            sh.fill.fore_color.rgb = BLUE
            sh.line.fill.background()
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.color.rgb = WHITE


def main():
    prs = Presentation(SRC)
    slides = list(prs.slides)

    for scene, (num, name, top_px) in HERO.items():
        slide = slides[num - 1]
        frame, label = by_id(slide, 14), by_id(slide, 15)
        add_picture(slide, name, frame.left, frame.top, frame.width, frame.height, top_px)
        remove(frame)
        remove(label)
    for num, shapes in TEXT.items():
        for shape_id, lines in shapes.items():
            set_paragraphs(by_id(slides[num - 1], shape_id), lines)

    # Scenes 6-8 are built now: vision badges become core demo badges.
    for num in (11, 12, 13):
        badge_core(slides[num - 1], 13, 1.60)
    overview = slides[2]
    for shape_id in (37, 43, 49):
        badge_core(overview, shape_id, 1.55)
    for shape_id in (63, 64):
        remove(by_id(overview, shape_id))
    for slide in slides:
        recolor_vision(slide)
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if VISION_NOTE.strip() in tf.text:
                tf.text = tf.text.replace(VISION_NOTE, "").replace(VISION_NOTE.strip(), "").rstrip()
            if tf.text.startswith("Vision scene. "):
                tf.text = tf.text[len("Vision scene. "):]
    for num, text in NOTES.items():
        slides[num - 1].notes_slide.notes_text_frame.text = text
    for num, (old, new) in NOTE_EDITS.items():
        tf = slides[num - 1].notes_slide.notes_text_frame
        tf.text = tf.text.replace(old, new)
    takeaways = slides[14]
    for shape_id in (14, 20):
        remove(by_id(takeaways, shape_id))

    table = next(s for s in slides[15].shapes if s.has_table).table
    for row in list(table.rows)[1:]:
        scene_cell, fallback = row.cells[1], row.cells[5]
        if "(vision)" in scene_cell.text:
            set_paragraphs(scene_cell, [scene_cell.text.replace("(vision)", "").strip()])
        if fallback.text == "Storyboard slide":
            set_paragraphs(fallback, ["Pre-built records (appendix screenshots)"])
        if fallback.text.startswith("Core demo only"):
            set_paragraphs(fallback, ["Act 1 only: 8.5 min"])

    # Appendix: every screenshot per scene, inserted before the closing slide.
    layout = slides[15].slide_layout
    closing = prs.slides._sldIdLst[-1]
    names = sorted(f[:-4] for f in os.listdir(SHOTS) if f.endswith(".png"))
    for scene in range(1, 10):
        shots = [n for n in names if n.startswith(f"s{scene}-")]
        slide = prs.slides.add_slide(layout)
        slide.placeholders[0].text = f"Scene {scene} · {TITLES[scene]}"
        slide.placeholders[1].text = "Screenshots from the fsl-rc org run"
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx > 1:
                remove(ph)
        cols = 3 if len(shots) > 4 else 2 if len(shots) > 1 else 1
        rows = math.ceil(len(shots) / cols)
        gap, cap = Inches(0.3), Inches(0.35)
        area_l, area_t, area_w, area_h = Inches(0.94), Inches(2.85), Inches(18.12), Inches(8.0)
        cell_w = (area_w - gap * (cols - 1)) / cols
        cell_h = min((area_h - gap * (rows - 1)) / rows - cap, cell_w / 1.68)
        for i, name in enumerate(shots):
            r, c = divmod(i, cols)
            left = area_l + c * (cell_w + gap)
            top = area_t + r * (cell_h + cap + gap)
            w, h = Image.open(os.path.join(SHOTS, name + ".png")).size
            top_px = APPENDIX_CROP.get(name, 0) if h > w else None
            add_picture(slide, name, int(left), int(top), int(cell_w), int(cell_h), top_px)
            box = slide.shapes.add_textbox(int(left), int(top + cell_h + Inches(0.05)), int(cell_w), int(cap))
            para = box.text_frame.paragraphs[0]
            run = para.add_run()
            run.text = CAPTIONS.get(name, name)
            run.font.size, run.font.name, run.font.color.rgb = Pt(13), "Salesforce Sans", GREY
        closing.addprevious(prs.slides._sldIdLst[-1])
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
